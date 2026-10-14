"""Project documents uploaded with a P&ID: piping specification and risk matrix, plus optional references.

The piping specification is read from a table with the columns Piping Specification | Services | ANSI Rating |
Material | Design Temperature | Design Pressure (multi-line temperature / pressure cells give the P-T rating).
The risk matrix is a grid with consequence levels down the side, likelihood levels across the top (or bottom) and
risk rankings (Low / Medium / High / Very High, or L / M / H / VH) in the cells; any size from 3 x 3 to 8 x 8.

Both can be an Excel workbook, a Word document (its tables) or a PDF / image of the table (read with OCR; ruled
tables work best). Excel and Word give exact values; OCR'd documents should be checked.
"""

import colorsys
import io
import re
from collections import Counter

import cv2
import docx
import numpy as np
import openpyxl
import pymupdf
from docx.oxml.ns import qn
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PIL import Image

from . import knowledge as kb
from . import risk
from .pipeline import ocr
from .pipeline.render import MAX_IMAGE_SIDE

SPREADSHEET_EXT = {".xlsx", ".xlsm"}
WORD_EXT = {".docx"}
IMAGE_EXT = {".pdf", ".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}
PROJECT_DOC_EXT = SPREADSHEET_EXT | WORD_EXT | IMAGE_EXT
OPTIONAL_EXT = IMAGE_EXT | WORD_EXT | SPREADSHEET_EXT | {".txt"}
MAX_OPTIONAL_DOCS = 5
MAX_TEXT_CHARS = 12000
OCR_DOC_SIDE = 1900  # the OCR engine shrinks anything over 2000 px
DOC_DPI = 200


class DocumentError(ValueError):
    pass


def ext_of(name):
    return "." + name.rsplit(".", 1)[-1].lower() if "." in (name or "") else ""


def _text(v):
    return "" if v is None else str(v).strip()


# ------------------------------------------------------------- tables from any document
# Word tables and OCR'd images are exposed through the small part of the openpyxl worksheet API the parsers use.

class _Cell:
    __slots__ = ("row", "column", "value", "color")

    def __init__(self, row, column, value=None, color=None):
        self.row, self.column, self.value, self.color = row, column, value, color


class _Table:
    def __init__(self, title, rows, lines=None):
        """rows: list of rows, each a list of (text, '#RRGGBB' or None); lines: the page's text lines, if known."""
        self.title = title
        self.lines = lines
        self.max_row = len(rows)
        self.max_column = max((len(r) for r in rows), default=0)
        self._cells = {(r, c): _Cell(r, c, text or None, color)
                       for r, row in enumerate(rows, 1) for c, (text, color) in enumerate(row, 1)}

    def cell(self, row, column):
        return self._cells.get((row, column)) or _Cell(row, column)

    def iter_rows(self, min_row=1, max_row=None, values_only=False):
        for r in range(min_row, (max_row or self.max_row) + 1):
            cells = tuple(self.cell(r, c) for c in range(1, self.max_column + 1))
            yield tuple(c.value for c in cells) if values_only else cells


class _Book:
    def __init__(self, tables):
        self.worksheets = tables


def _word_tables(data):
    d = docx.Document(io.BytesIO(data))
    tables, heading = [], ""
    for el in d.element.body.iterchildren():
        if el.tag == qn("w:p"):
            text = "".join(t.text or "" for t in el.iter(qn("w:t"))).strip()
            if text:
                heading = text
        elif el.tag == qn("w:tbl"):
            rows = []
            for row in docx.table.Table(el, d).rows:
                cells = []
                for cell in row.cells:
                    shd = cell._tc.tcPr.find(qn("w:shd")) if cell._tc.tcPr is not None else None
                    fill = shd.get(qn("w:fill")) if shd is not None else None
                    color = f"#{fill.upper()}" if fill and re.fullmatch(r"[0-9A-Fa-f]{6}", fill) and fill.upper() not in ("FFFFFF", "000000") else None
                    cells.append((cell.text.strip(), color))
                rows.append(cells)
            tables.append(_Table(heading or f"Table {len(tables) + 1}", rows))
    return tables


def _image(data, filename):
    """RGB array of the document (first page for PDFs) scaled for OCR, plus any PDF words in its pixels."""
    if ext_of(filename) == ".pdf":
        doc = pymupdf.open(stream=data, filetype="pdf")
        if doc.page_count == 0:
            raise DocumentError(f"'{filename}' has no pages.")
        page = doc[0]
        dpi = min(DOC_DPI, int(MAX_IMAGE_SIDE * 72 / max(page.rect.width, page.rect.height)))
        pix = page.get_pixmap(dpi=dpi, colorspace=pymupdf.csRGB, alpha=False)
        rgb = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width, 3).copy()
        s = dpi / 72
        words = [{"text": w[4], "box": [w[0] * s, w[1] * s, w[2] * s, w[3] * s]} for w in page.get_text("words")]
        return rgb, words
    img = Image.open(io.BytesIO(data)).convert("RGB")
    scale = min(MAX_IMAGE_SIDE, max(OCR_DOC_SIDE, max(img.size))) / max(img.size)
    if abs(scale - 1) > 0.01:
        img = img.resize((round(img.width * scale), round(img.height * scale)), Image.LANCZOS)
    return np.array(img), []


def _line_positions(mask, horizontal, min_len):
    """Centres of the ruling lines in a line mask: unbroken segments at least min_len long (letter strokes are shorter)."""
    n, _, stats, _ = cv2.connectedComponentsWithStats(mask)
    length, pos, extent = (2, 1, 3) if horizontal else (3, 0, 2)
    centres = sorted(stats[i, pos] + stats[i, extent] / 2 for i in range(1, n) if stats[i, length] >= min_len)
    out = []
    for c in centres:
        if out and c - out[-1][-1] <= 6:
            out[-1].append(c)
        else:
            out.append([c])
    return [int(np.mean(g)) for g in out]


def _cell_color(rgb, x0, y0, x1, y1):
    pad_x, pad_y = max(2, (x1 - x0) // 8), max(2, (y1 - y0) // 8)
    crop = rgb[y0 + pad_y:y1 - pad_y, x0 + pad_x:x1 - pad_x].reshape(-1, 3)
    if not len(crop):
        return None
    bright = crop[crop.max(axis=1) > 90]
    if not len(bright):
        return None
    q = (bright // 16).astype(np.int32)
    keys, counts = np.unique(q[:, 0] * 256 + q[:, 1] * 16 + q[:, 2], return_counts=True)
    k = keys[counts.argmax()]
    sel = bright[(q[:, 0] * 256 + q[:, 1] * 16 + q[:, 2]) == k]
    r, g, b = (int(v) for v in np.median(sel, axis=0))
    if min(r, g, b) > 235 or max(r, g, b) - min(r, g, b) < 18:
        return None
    return f"#{r:02X}{g:02X}{b:02X}"


def _merge_words(tokens):
    """Join word-level tokens (PDF text) that sit next to each other on a line into phrases."""
    out = []
    for t in sorted(tokens, key=lambda t: ((t["box"][1] + t["box"][3]) / 2, t["box"][0])):
        x0, y0, x1, y1 = t["box"]
        h = y1 - y0
        last = out[-1] if out else None
        if last and abs((last["box"][1] + last["box"][3]) / 2 - (y0 + y1) / 2) < 0.5 * h \
                and 0 <= x0 - last["box"][2] < 0.8 * h:
            last["text"] += " " + t["text"]
            lb = last["box"]
            last["box"] = [lb[0], min(lb[1], y0), x1, max(lb[3], y1)]
        else:
            out.append({"text": t["text"], "box": [x0, y0, x1, y1]})
    return out


def _ruled_table(rgb, tokens):
    if not tokens:
        return None
    text_h = float(np.median([t["box"][3] - t["box"][1] for t in tokens]))
    gray = cv2.cvtColor(rgb, cv2.COLOR_RGB2GRAY)
    bw = cv2.adaptiveThreshold(gray, 255, cv2.ADAPTIVE_THRESH_MEAN_C, cv2.THRESH_BINARY_INV, 25, 12)
    H, W = gray.shape
    horiz = cv2.morphologyEx(bw, cv2.MORPH_OPEN, cv2.getStructuringElement(cv2.MORPH_RECT, (max(25, W // 30), 1)))
    vert = cv2.morphologyEx(bw, cv2.MORPH_OPEN, cv2.getStructuringElement(cv2.MORPH_RECT, (1, max(25, H // 30))))
    ys = _line_positions(horiz, True, W * 0.25)
    xs = _line_positions(vert, False, max(H * 0.05, 1.6 * text_h))
    if len(ys) < 3 or len(xs) < 3:
        return None
    rows = [[[] for _ in range(len(xs) - 1)] for _ in range(len(ys) - 1)]
    for t in tokens:
        cx, cy = (t["box"][0] + t["box"][2]) / 2, (t["box"][1] + t["box"][3]) / 2
        r = next((i for i in range(len(ys) - 1) if ys[i] <= cy < ys[i + 1]), None)
        c = next((j for j in range(len(xs) - 1) if xs[j] <= cx < xs[j + 1]), None)
        if r is not None and c is not None:
            rows[r][c].append(t)
    out = []
    for i, row in enumerate(rows):
        cells = []
        for j, toks in enumerate(row):
            text = "\n".join(t["text"] for t in _merge_words(toks))
            cells.append((text, _cell_color(rgb, xs[j], ys[i], xs[j + 1], ys[i + 1])))
        if any(text for text, _ in cells):
            out.append(cells)
    return out


def _token_table(tokens):
    """Rows and columns from text positions alone (tables without ruling lines)."""
    if not tokens:
        return []
    h = float(np.median([t["box"][3] - t["box"][1] for t in tokens]))
    lines = []
    for t in sorted(tokens, key=lambda t: (t["box"][1] + t["box"][3]) / 2):
        cy = (t["box"][1] + t["box"][3]) / 2
        if lines and abs(lines[-1]["cy"] - cy) < 0.6 * h:
            lines[-1]["toks"].append(t)
        else:
            lines.append({"cy": cy, "toks": [t]})
    spans = sorted((t["box"][0], t["box"][2]) for line in lines if len(line["toks"]) > 1 for t in line["toks"])
    bounds, end = [], None
    for x0, x1 in spans:
        if end is not None and x0 - end > 0.8 * h:
            bounds.append((end + x0) / 2)
        end = x1 if end is None else max(end, x1)
    col = lambda x: sum(1 for b in bounds if x >= b)
    rows = []
    for line in lines:
        cells = [[] for _ in range(len(bounds) + 1)]
        for t in sorted(line["toks"], key=lambda t: t["box"][0]):
            cells[col((t["box"][0] + t["box"][2]) / 2)].append(t["text"])
        texts = [" ".join(c) for c in cells]
        filled = sum(1 for x in texts if x)
        prev = rows[-1] if rows else None
        if prev and not texts[0] and prev[0] and filled < sum(1 for x in prev if x):
            for k, x in enumerate(texts):
                if x:
                    prev[k] = f"{prev[k]}\n{x}" if prev[k] else x
        else:
            rows.append(texts)
    return [[(text, None) for text in row] for row in rows]


def _image_tables(data, filename):
    rgb, words = _image(data, filename)
    tokens = [t for t in (words or ocr.read_document(rgb)) if t["text"].strip()]
    lines = _merge_words(tokens)
    rows = _ruled_table(rgb, tokens) or _token_table(lines)
    return [_Table(filename, rows, [t["text"] for t in lines])] if rows else []


def _scan_hint(filename):
    if ext_of(filename) in IMAGE_EXT:
        return " For an image or scanned PDF use a clear, straight, high-resolution copy of the table, or upload Excel / Word."
    return ""


def _tables(data, filename, label):
    """Open an Excel workbook, Word document, PDF or image as a book of tables."""
    ext = ext_of(filename)
    try:
        if ext in SPREADSHEET_EXT:
            return openpyxl.load_workbook(io.BytesIO(data), data_only=True)
        if ext in WORD_EXT:
            tables = _word_tables(data)
            if not tables:
                raise DocumentError(f"'{filename}' has no table. Put the {label} in a Word table.")
            return _Book(tables)
        if ext in IMAGE_EXT:
            return _Book(_image_tables(data, filename))
    except DocumentError:
        raise
    except Exception as e:
        raise DocumentError(f"The {label} '{filename}' could not be read: {e}") from e
    raise DocumentError(f"Unsupported {label} file type '{ext or filename}'. Use Excel, Word, PDF or an image.")


# ------------------------------------------------------------------ piping specification

SPEC_COLUMNS = {
    "spec": re.compile(r"spec", re.I),
    "services": re.compile(r"service", re.I),
    "rating": re.compile(r"rating|class", re.I),
    "material": re.compile(r"material", re.I),
    "temp": re.compile(r"temp", re.I),
    "pressure": re.compile(r"press", re.I),
}
NUM = re.compile(r"-?\d+(?:\.\d+)?")


def _spec_header(ws):
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 30)):
        found = {}
        for cell in row:
            t = _text(cell.value)
            if not t:
                continue
            for key, rx in SPEC_COLUMNS.items():
                if key not in found and rx.search(t) and not (key == "spec" and SPEC_COLUMNS["services"].search(t)):
                    found[key] = cell.column
                    break
        if {"spec", "rating", "temp", "pressure"} <= found.keys():
            return row[0].row, found
    return None, None


def _temps_f(text):
    vals = [float(x) for x in NUM.findall(text)]
    if re.search(r"\d\s*\u00b0?\s*C\b", text, re.I) and not re.search(r"\d\s*\u00b0?\s*F\b", text, re.I):
        vals = [v * 9 / 5 + 32 for v in vals]
    return [round(v) for v in vals]


def _pressures_psig(text):
    vals = [float(x) for x in NUM.findall(text)]
    low = text.lower()
    factor = 14.5038 if "bar" in low else 0.145038 if "kpa" in low else 145.038 if "mpa" in low else 1
    return [round(v * factor) for v in vals]


def _rating(text):
    up = re.sub(r"(?<=\d)O+(?![A-Z])", lambda m: "0" * len(m.group()), text.upper())  # OCR reads 0 as O
    up = re.sub(r"(?<=\d)\s+(?=\d)", "", up)
    m = re.search(r"CLASS\s*(\d+)", up) or re.search(r"(\d+)\s*(?:#|LB)", up) or re.search(r"\b(\d{3,4})\b", up)
    return f"Class {m.group(1)}" if m else (text.split(",")[0].strip() or None)


def parse_piping_spec(data, filename="piping specification"):
    """Return {spec_code: spec} in the knowledge-base schema."""
    wb = _tables(data, filename, "piping specification")
    service_codes = set(kb.services()["codes"])
    problems = []
    for ws in wb.worksheets:
        header_row, cols = _spec_header(ws)
        if not header_row:
            continue
        specs = {}
        for r in range(header_row + 1, ws.max_row + 1):
            get = lambda key: _text(ws.cell(row=r, column=cols[key]).value) if key in cols else ""
            code_text = get("spec")
            if not code_text:
                continue
            code = re.split(r"[\s,;/]+", code_text.upper())[0].strip(".:-")
            if not code:
                continue
            temps, pressures = _temps_f(get("temp")), _pressures_psig(get("pressure"))
            if not pressures or not temps:
                problems.append(f"{code}: design temperature / pressure not readable")
                continue
            if len(temps) == len(pressures) + 1:
                points_t = temps[1:]
            elif len(temps) == len(pressures):
                points_t = temps
            elif len(pressures) == 1:
                points_t = [temps[-1]]
            else:
                problems.append(f"{code}: {len(temps)} temperatures but {len(pressures)} pressures")
                continue
            rating_text = get("rating")
            services = [s for s in re.findall(r"\b[A-Z]{1,3}\b", get("services").upper()) if s in service_codes]
            specs[code] = {
                "rating": _rating(rating_text),
                "sour_service": "SOUR" in (rating_text + " " + get("material") + " " + get("services")).upper(),
                "services": list(dict.fromkeys(services)),
                "material": re.sub(r"\s+", " ", get("material")) or None,
                "min_temp_f": temps[0] if len(temps) > 1 else None,
                "max_temp_f": temps[-1],
                "pressure_temp": [{"temp_f": t, "psig": p} for t, p in zip(points_t, pressures)],
            }
        if specs:
            return specs
    detail = f" Problems: {'; '.join(problems)}." if problems else ""
    raise DocumentError(
        f"No piping specification table was found in '{filename}'. Use a table with the columns Piping Specification, "
        f"Services, ANSI Rating, Material, Design Temperature and Design Pressure.{detail}{_scan_hint(filename)}")


def spec_summary(specs):
    return [{"spec": code, "rating": s["rating"], "sour_service": s["sour_service"], "min_temp_f": s["min_temp_f"],
             "max_temp_f": s["max_temp_f"], "psig_at_min": s["pressure_temp"][0]["psig"]} for code, s in specs.items()]


# ----------------------------------------------------------------------- risk matrix

LABELS = {
    "L": ("L", "LOW", "VERY LOW", "VL", "NEGLIGIBLE", "ACCEPTABLE", "BROADLY ACCEPTABLE", "GREEN"),
    "M": ("M", "MED", "MEDIUM", "MODERATE", "ALARP", "TOLERABLE", "TOLERABLE IF ALARP", "YELLOW"),
    "H": ("H", "HIGH", "SIGNIFICANT", "SERIOUS", "AMBER", "ORANGE"),
    "VH": ("VH", "VERY HIGH", "EXTREME", "CRITICAL", "INTOLERABLE", "UNACCEPTABLE", "SEVERE", "RED"),
}
LABEL_CODE = {label.replace(" ", ""): code for code, labels in LABELS.items() for label in labels}
LEVEL_RE = re.compile(r"^\s*(?:[A-Za-z]{1,2}\s*)?(\d{1,2})(?:\s*[-\u2013:.)]\s*(.*)|\s+([A-Za-z].*))?\s*$", re.S)
LIKELIHOOD_WORDS = re.compile(r"likelihood|probability|frequency", re.I)
CONSEQUENCE_WORDS = re.compile(r"consequence|severity|impact", re.I)


def _level(v):
    if isinstance(v, bool):
        return None, ""
    if isinstance(v, (int, float)) and float(v).is_integer():
        return int(v), ""
    m = LEVEL_RE.match(_text(v))
    return (int(m.group(1)), (m.group(2) or m.group(3) or "").strip()) if m else (None, "")


def _label_code(v):
    t = re.sub(r"[^A-Z ]+", " ", _text(v).upper())
    t = re.sub(r"\bRISK\b", " ", t)
    t = re.sub(r"\s+", " ", t).strip()
    if t.replace(" ", "") in LABEL_CODE:
        return LABEL_CODE[t.replace(" ", "")]
    for token in ("VH", "H", "M", "L"):
        if re.search(rf"\b{token}\b", t):
            return token
    return None


def _fill_hex(cell):
    if isinstance(cell, _Cell):
        return cell.color
    fill = cell.fill
    if fill is None or fill.fill_type != "solid":
        return None
    rgb = getattr(fill.fgColor, "rgb", None)
    if fill.fgColor.type != "rgb" or not isinstance(rgb, str) or len(rgb) != 8 or rgb[2:] in ("000000",):
        return None
    return "#" + rgb[2:].upper()


def _run(ws, r, c, dr, dc, limit=risk.MAX_LEVELS):
    """Consecutive level numbers starting at (r, c) moving by (dr, dc)."""
    out = []
    while len(out) < limit + 1:
        if r < 1 or c < 1 or r > ws.max_row or c > ws.max_column:
            break
        n, name = _level(ws.cell(row=r, column=c).value)
        if n is None:
            break
        out.append((n, name, r, c))
        r, c = r + dr, c + dc
    nums = [x[0] for x in out]
    return out if risk.MIN_LEVELS <= len(nums) <= risk.MAX_LEVELS and sorted(nums) == list(range(1, len(nums) + 1)) else None


def _find_grid(ws):
    for r in range(1, min(ws.max_row, 60) + 1):
        for c in range(1, min(ws.max_column, 25) + 1):
            cols = _run(ws, r, c + 1, 0, 1)
            if not cols:
                continue
            for dr in (1, -1):
                rows = _run(ws, r + dr, c, dr, 0)
                if rows:
                    corner = _text(ws.cell(row=r, column=c).value)
                    return rows, cols, corner
    return None


def _level_sheet(wb, words):
    for ws in wb.worksheets:
        if not words.search(ws.title):
            continue
        out = {}
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 20), values_only=True):
            vals = [v for v in row if v not in (None, "")]
            if len(vals) >= 2:
                n, _ = _level(vals[0])
                if n:
                    out[n] = {"name": _text(vals[1]), "desc": _text(vals[2]) if len(vals) > 2 else ""}
        if out:
            return out
    return {}


# Matrices whose cells hold risk numbers (1, 2 ... 10 or C x L scores) and / or colours, with the level indices anywhere
# beside the grid (e.g. likelihood 1 = Likely down the side, consequence 6 ... 1 in a row under the grid).

LEGEND_RE = re.compile(r"^\s*(\d{1,2}(?:\s*(?:[,.;&/]|and)\s*\d{1,2})*)\s*[-\u2013:)=]\s*([A-Za-z(].*)$", re.I | re.S)
WORST_WORDS = re.compile(r"intolerable|unacceptable|immediate|short[- ]?term|stop|not permitted|extreme", re.I)
BEST_WORDS = re.compile(r"\bacceptable|no further|manage risk|broadly|negligible|routine", re.I)
CODES_BY_COUNT = {1: ["M"], 2: ["H", "L"], 3: ["H", "M", "L"], 4: ["VH", "H", "M", "L"]}
BAND_ORDER_WORST = ["VH", "H", "M", "L"]
BAND_HUE = {"VH": 0, "H": 30, "M": 60, "L": 100}


def _hue(color):
    """Hue in degrees (red ~0, yellow 60, green 120) or None for grey / missing colours."""
    if not color:
        return None
    r, g, b = (int(color[k:k + 2], 16) / 255 for k in (1, 3, 5))
    h, s, _ = colorsys.rgb_to_hsv(r, g, b)
    if s < 0.15:
        return None
    h *= 360
    return h - 360 if h > 300 else h


def _hue_code(h):
    return "VH" if h <= 12 else "H" if h <= 50 else "M" if h <= 70 else "L"


def _rgb_dist(a, b):
    return sum((int(a[k:k + 2], 16) - int(b[k:k + 2], 16)) ** 2 for k in (1, 3, 5)) ** 0.5


def _cell_value(v):
    if isinstance(v, (int, float)) and not isinstance(v, bool) and float(v).is_integer():
        return int(v)
    t = _text(v)
    if re.fullmatch(r"\d{1,3}", t):
        return int(t)
    if 0 < len(t) <= 20 and re.search(r"[A-Za-z]", t):
        return _label_code(t)
    return None


def _value_block(ws):
    """(r0, r1, c0, c1) of the largest block of risk values / coloured blanks, 3 x 3 to 8 x 8."""
    R, C = min(ws.max_row, 80), min(ws.max_column, 40)
    vals, run = {}, {}
    for r in range(1, R + 1):
        for c in range(C, 0, -1):
            cell = ws.cell(row=r, column=c)
            t = _text(cell.value)
            vals[r, c] = _cell_value(cell.value)
            filled = vals[r, c] is not None or (not t and _fill_hex(cell) is not None)
            run[r, c] = run.get((r, c + 1), 0) + 1 if filled else 0
    lo, hi = risk.MIN_LEVELS, risk.MAX_LEVELS
    best = None
    for r0 in range(1, R + 1):
        for c0 in range(1, C + 1):
            w = hi
            for r1 in range(r0, min(R, r0 + hi - 1) + 1):
                w = min(w, run[r1, c0])
                if w < lo:
                    break
                h = r1 - r0 + 1
                if h < lo:
                    continue
                n = sum(1 for r in range(r0, r1 + 1) for c in range(c0, c0 + w) if vals[r, c] is not None)
                if n >= 0.75 * h * w and (best is None or (h * w, n) > best[0]):
                    best = ((h * w, n), [r0, r1, c0, c0 + w - 1])
    if not best:
        return None
    r0, r1, c0, c1 = best[1]

    def axis_like(cells):
        nums = [vals[rc] for rc in cells]
        colors = {_fill_hex(ws.cell(row=r, column=c)) for r, c in cells}
        return (all(isinstance(n, int) for n in nums) and len(colors) == 1
                and all(abs(b - a) == 1 for a, b in zip(nums, nums[1:])) and len(set(nums)) == len(nums))

    inner_colors = {_fill_hex(ws.cell(row=r, column=c)) for r in range(r0, r1 + 1) for c in range(c0, c1 + 1)}
    while len(inner_colors - {None}) >= 2:
        if r1 - r0 + 1 > lo and axis_like([(r0, c) for c in range(c0, c1 + 1)]):
            r0 += 1
        elif r1 - r0 + 1 > lo and axis_like([(r1, c) for c in range(c0, c1 + 1)]):
            r1 -= 1
        elif c1 - c0 + 1 > lo and axis_like([(r, c0) for r in range(r0, r1 + 1)]):
            c0 += 1
        elif c1 - c0 + 1 > lo and axis_like([(r, c1) for r in range(r0, r1 + 1)]):
            c1 -= 1
        else:
            break
    return r0, r1, c0, c1


def _legend_groups(lines, nums):
    """Groups of risk numbers from legend lines such as '1, 2, 3, 4 - Short-term risk reduction required'."""
    cands = []
    for line in lines:
        for part in str(line).split("\n"):
            m = LEGEND_RE.match(part)
            if m:
                g = [int(x) for x in re.findall(r"\d{1,2}", m.group(1))]
                if set(g) & nums and max(g) <= max(nums) + 1:
                    cands.append((g, m.group(2)))
    groups, used = [], set()
    for g, text in sorted(cands, key=lambda x: -len(x[0])):
        if not used & set(g):
            groups.append((g, text))
            used |= set(g)
    return [(g, t) for g, t in groups if set(g) & nums] if nums <= used and len(groups) >= 2 else None


def _axis_names(ws, positions, across, outside):
    """Level names beside the grid: the nearest row / column outside the grid with short text for most levels."""
    for k in outside:
        texts = [_text(ws.cell(row=k, column=p).value if across else ws.cell(row=p, column=k).value) for p in positions]
        names = [re.sub(r"^\s*\d{1,2}\s*[-\u2013:.)]?\s*", "", t) for t in texts]
        good = [n for n in names if n and len(n) <= 40 and re.search(r"[A-Za-z]{3}", n)]
        if len(good) >= 0.8 * len(positions):
            return names
    return [""] * len(positions)


def _score_grid(ws, filename):
    block = _value_block(ws)
    if not block:
        return None
    r0, r1, c0, c1 = block
    rows, cols = list(range(r0, r1 + 1)), list(range(c0, c1 + 1))
    cell = lambda r, c: ws.cell(row=r, column=c)
    vals = {(r, c): _cell_value(cell(r, c).value) for r in rows for c in cols}
    colors = {(r, c): _fill_hex(cell(r, c)) for r in rows for c in cols}
    known = {rc: v for rc, v in vals.items() if v is not None}
    for rc, v in vals.items():
        if v is None:
            near = [(_rgb_dist(colors[rc], colors[k]), kv) for k, kv in known.items() if colors[rc] and colors[k]]
            if not near or min(near)[0] > 40:
                return None
            vals[rc] = min(near)[1]

    if all(isinstance(v, str) for v in vals.values()):
        codes = dict(vals)
    elif all(isinstance(v, int) for v in vals.values()):
        nums = set(vals.values())
        hue_of = {}
        for n in nums:
            hues = [_hue(colors[rc]) for rc, v in vals.items() if v == n and _hue(colors[rc]) is not None]
            hue_of[n] = float(np.median(hues)) if hues else None
        lines = ws.lines if getattr(ws, "lines", None) else [
            _text(c.value) for row in ws.iter_rows() for c in row if _text(c.value)]
        groups = _legend_groups(lines, nums)
        band_of = {}
        if groups:
            def worse_first(item):
                g, text = item
                hues = [hue_of[n] for n in g if hue_of.get(n) is not None]
                if len(hues) == len(set(g) & nums):
                    return (0, float(np.mean(hues)))
                words = len(BEST_WORDS.findall(text)) - len(WORST_WORDS.findall(text))
                return (1, words, -max(g))
            groups.sort(key=worse_first)
            all_hues = all(hue_of[n] is not None for n in nums)
            for i, (g, _) in enumerate(groups):
                if len(groups) in CODES_BY_COUNT:
                    code = CODES_BY_COUNT[len(groups)][i]
                elif all_hues:
                    code = _hue_code(float(np.mean([hue_of[n] for n in g if n in nums])))
                else:
                    code = BAND_ORDER_WORST[round(i * 3 / (len(groups) - 1))]
                band_of.update({n: code for n in g})
        elif all(h is not None for h in hue_of.values()):
            band_of = {n: _hue_code(h) for n, h in hue_of.items()}
        else:
            return None
        codes = {rc: band_of[v] for rc, v in vals.items()}
    else:
        return None

    rank = {rc: risk.BAND_ORDER.index(code) for rc, code in codes.items()}
    mean_rank = np.mean(list(rank.values()))

    def trend(index_of):
        pos = {rc: index_of(rc) for rc in rank}
        mean_pos = np.mean(list(pos.values()))
        return sum((pos[rc] - mean_pos) * (rank[rc] - mean_rank) for rc in rank)

    row_trend, col_trend = trend(lambda rc: rc[0]), trend(lambda rc: rc[1])
    # Levels count up towards the worse end of each axis; flat axes keep the usual layout (highest row at the top).
    row_level = {r: (i + 1 if row_trend > 0 else len(rows) - i) for i, r in enumerate(rows)}
    col_level = {c: (len(cols) - i if col_trend < 0 else i + 1) for i, c in enumerate(cols)}

    side_cols = sorted((c for c in range(1, ws.max_column + 1) if c not in cols), key=lambda c: min(abs(c - c0), abs(c - c1)))
    side_rows = sorted((r for r in range(1, ws.max_row + 1) if r not in rows), key=lambda r: min(abs(r - r0), abs(r - r1)))
    row_names = _axis_names(ws, rows, False, side_cols)
    col_names = _axis_names(ws, cols, True, side_rows)

    # Which axis is likelihood: count the axis words in the headings beside the grid rows and around the grid columns.
    def words(cells):
        text = " ".join(_text(cell(r, c).value) for r, c in cells)
        return len(LIKELIHOOD_WORDS.findall(text)) - len(CONSEQUENCE_WORDS.findall(text))

    beside_rows = [(r, c) for r in range(1, r1 + 1) for c in side_cols]
    below = [r for r in side_rows if r1 < r <= r1 + 4]
    above = [r for r in side_rows if r < r0]
    beside_cols = [(r, c) for r in below for c in range(1, ws.max_column + 1)] + [(r, c) for r in above for c in cols]
    rows_are_likelihood = words(beside_rows) > words(beside_cols)

    row_axis = sorted(({"level": row_level[r], "name": n} for r, n in zip(rows, row_names)), key=lambda x: x["level"])
    col_axis = sorted(({"level": col_level[c], "name": n} for c, n in zip(cols, col_names)), key=lambda x: x["level"])
    sev_axis, lik_axis = (col_axis, row_axis) if rows_are_likelihood else (row_axis, col_axis)
    nc, nl = len(sev_axis), len(lik_axis)
    cells = [[None] * nl for _ in range(nc)]
    for (r, c), code in codes.items():
        s, l = (col_level[c], row_level[r]) if rows_are_likelihood else (row_level[r], col_level[c])
        cells[s - 1][l - 1] = code

    used = sorted(set(codes.values()), key=risk.BAND_ORDER.index)

    def band_color(code):
        """The band's cell colour closest to its usual hue (graded matrices have several shades per band)."""
        shades = {colors[rc] for rc, x in codes.items() if x == code and _hue(colors[rc]) is not None}
        if not shades:
            return risk.BAND_DEFAULTS[code]["color"]
        return min(shades, key=lambda s: (abs(_hue(s) - BAND_HUE[code]), s))
    generic_sev = risk.SEVERITY if nc == 5 else []
    generic_lik = risk.LIKELIHOOD if nl == 5 else []
    level = lambda it, generic: {"level": it["level"],
                                 "name": it["name"] or (generic[it["level"] - 1]["name"] if generic else f"Level {it['level']}"),
                                 "desc": ""}
    return risk.validate({
        "name": f"{nc} x {nl} project risk matrix",
        "source": filename,
        "severity": [level(it, generic_sev) for it in sev_axis],
        "likelihood": [level(it, generic_lik) for it in lik_axis],
        "bands": [{"code": code, "name": risk.BAND_DEFAULTS[code]["name"], "color": band_color(code)} for code in used],
        "cells": cells,
    })


def parse_risk_matrix(data, filename="risk matrix"):
    wb = _tables(data, filename, "risk matrix")
    sheets = sorted(wb.worksheets, key=lambda s: 0 if re.search(r"matrix", s.title, re.I) else 1)
    problem = None
    for ws in sheets:
        try:
            matrix = _label_grid(wb, ws, filename)
        except DocumentError as e:
            problem, matrix = problem or e, None
        matrix = matrix or _score_grid(ws, filename)
        if matrix:
            return kb.apply_matrix_definitions(matrix)
    if problem:
        raise problem
    raise DocumentError(
        f"No risk matrix grid was found in '{filename}'. Put consequence levels (1, 2, 3...) down the first column, "
        f"likelihood levels across the header row and Low / Medium / High / Very High (or risk numbers with a legend) "
        f"in the cells.{_scan_hint(filename)}")


def _label_grid(wb, ws, filename):
    """Grid with level numbers on its first row / column and Low ... Very High in the cells; None if not found."""
    grid = _find_grid(ws)
    if not grid:
        return None
    rows, cols, corner = grid
    transposed = bool(LIKELIHOOD_WORDS.search(corner)) and (
        not CONSEQUENCE_WORDS.search(corner)
        or LIKELIHOOD_WORDS.search(corner).start() < CONSEQUENCE_WORDS.search(corner).start())
    cells, unknown, colors, names = {}, set(), {}, {}
    for n_row, _, r, _ in rows:
        for n_col, _, _, c in cols:
            cell = ws.cell(row=r, column=c)
            code = _label_code(cell.value)
            if not code:
                unknown.add(_text(cell.value) or "(blank)")
                continue
            key = (n_col, n_row) if transposed else (n_row, n_col)
            cells[key] = code
            fill = _fill_hex(cell)
            if fill:
                colors.setdefault(code, Counter())[fill] += 1
            label = _text(cell.value)
            if label.replace(" ", "").upper() == risk.BAND_DEFAULTS[code]["name"].replace(" ", "").upper():
                label = risk.BAND_DEFAULTS[code]["name"]
            if len(label) > 2 and not re.search(r"\d", label):
                names.setdefault(code, Counter())[label.title() if label.isupper() and len(label) > 5 else label] += 1
    if unknown:
        raise DocumentError(
            f"The risk matrix in '{filename}' has cells that are not a risk ranking: {', '.join(sorted(unknown)[:6])}. "
            f"Use Low, Medium, High or Very High (or L, M, H, VH) in every cell, or risk numbers with a legend."
            f"{_scan_hint(filename)}")
    sev_axis, lik_axis = (cols, rows) if transposed else (rows, cols)
    nc, nl = len(sev_axis), len(lik_axis)
    sev_names = {n: name for n, name, _, _ in sev_axis}
    lik_names = {n: name for n, name, _, _ in lik_axis}
    sev_sheet = _level_sheet(wb, CONSEQUENCE_WORDS)
    lik_sheet = _level_sheet(wb, LIKELIHOOD_WORDS)
    generic_sev = risk.SEVERITY if nc == 5 else []
    generic_lik = risk.LIKELIHOOD if nl == 5 else []

    def level(i, axis_names, sheet, generic):
        s = sheet.get(i, {})
        g = generic[i - 1] if generic else {}
        return {"level": i, "name": s.get("name") or axis_names.get(i) or g.get("name") or f"Level {i}",
                "desc": s.get("desc") or (g.get("desc", "") if not (s or axis_names.get(i)) else "")}

    used = sorted(set(cells.values()), key=risk.BAND_ORDER.index)
    return risk.validate({
        "name": f"{nc} x {nl} project risk matrix",
        "source": filename,
        "severity": [level(i, sev_names, sev_sheet, generic_sev) for i in range(1, nc + 1)],
        "likelihood": [level(i, lik_names, lik_sheet, generic_lik) for i in range(1, nl + 1)],
        "bands": [{"code": code,
                   "name": names[code].most_common(1)[0][0] if code in names else risk.BAND_DEFAULTS[code]["name"],
                   "color": colors[code].most_common(1)[0][0] if code in colors else risk.BAND_DEFAULTS[code]["color"]}
                  for code in used],
        "cells": [[cells[(c, l)] for l in range(1, nl + 1)] for c in range(1, nc + 1)],
    })


# ------------------------------------------------------------------------- templates

THIN = Side(style="thin", color="7F8C9D")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD = PatternFill("solid", fgColor="C0E0FF")
BOLD = Font(name="Arial", bold=True, size=10)
PLAIN = Font(name="Arial", size=10)
WRAP = Alignment(wrap_text=True, vertical="center")
CENTER = Alignment(wrap_text=True, vertical="center", horizontal="center")


def _put(ws, r, c, value, font=PLAIN, fill=None, align=WRAP):
    cell = ws.cell(row=r, column=c, value=value)
    cell.font, cell.alignment, cell.border = font, align, BORDER
    if fill:
        cell.fill = fill
    return cell


def _save(wb):
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def risk_matrix_template():
    m = risk.DEFAULT_MATRIX
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Risk Matrix"
    _put(ws, 1, 1, "Consequence \\ Likelihood", BOLD, HEAD, CENTER)
    for j, lk in enumerate(m["likelihood"], 2):
        _put(ws, 1, j, f"{lk['level']} - {lk['name']}", BOLD, HEAD, CENTER)
    for i, sev in enumerate(reversed(m["severity"]), 2):
        _put(ws, i, 1, f"{sev['level']} - {sev['name']}", BOLD, HEAD)
        for j, lk in enumerate(m["likelihood"], 2):
            band = risk.band_info(m["cells"][sev["level"] - 1][lk["level"] - 1], m)
            _put(ws, i, j, band["name"], BOLD, PatternFill("solid", fgColor=band["color"].lstrip("#")), CENTER)
        ws.row_dimensions[i].height = 30
    ws.column_dimensions["A"].width = 26
    for j in range(2, 2 + len(m["likelihood"]) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 16
    notes = [
        "How to use this template",
        "Rows are consequence levels (highest at the top), columns are likelihood levels. Keep the numbers 1, 2, 3... in the "
        "first column and header row.",
        "For a 6 x 6 matrix add a row '6 - ...' and a column '6 - ...'; any size from 3 x 3 to 8 x 8 is accepted.",
        "Every cell must say Low, Medium, High or Very High (or L, M, H, VH). Cell colours are used in the worksheet.",
        "Describe each level on the 'Consequence' and 'Likelihood' sheets.",
    ]
    for k, text in enumerate(notes, 9):
        ws.cell(row=k, column=1, value=text).font = BOLD if k == 9 else PLAIN
    for title, items in (("Consequence", m["severity"]), ("Likelihood", m["likelihood"])):
        s = wb.create_sheet(title)
        for j, h in enumerate(("Level", "Name", "Description"), 1):
            _put(s, 1, j, h, BOLD, HEAD, CENTER)
        for i, it in enumerate(items, 2):
            _put(s, i, 1, it["level"], align=CENTER)
            _put(s, i, 2, it["name"])
            _put(s, i, 3, it["desc"])
        s.column_dimensions["A"].width, s.column_dimensions["B"].width, s.column_dimensions["C"].width = 8, 18, 80
    return _save(wb)


def piping_spec_template():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Spec"
    header = ["Piping Specification", "Services", "ANSI Rating", "Material", "Design Temperature", "Design Pressure"]
    for j, h in enumerate(header, 1):
        _put(ws, 1, j, h, BOLD, HEAD, CENTER)
    services = kb.services()["codes"]
    for i, (code, s) in enumerate(kb.piping()["specs"].items(), 2):
        pts = s["pressure_temp"]
        temps = ([f"{s['min_temp_f']}F to {pts[0]['temp_f']}F"] if s.get("min_temp_f") is not None else [f"{pts[0]['temp_f']}F"]) \
            + [f"{p['temp_f']}F" for p in pts[1:]]
        _put(ws, i, 1, code, BOLD)
        _put(ws, i, 2, ", ".join(x for x in s.get("services", []) if x in services))
        _put(ws, i, 3, s["rating"] + (", Sour Service" if s.get("sour_service") else ""))
        _put(ws, i, 4, s.get("material") or "")
        _put(ws, i, 5, "\n".join(temps))
        _put(ws, i, 6, "\n".join(f"{p['psig']}psig" for p in pts))
        ws.row_dimensions[i].height = 15 * max(len(pts), 2)
    for j, w in enumerate((20, 18, 24, 48, 22, 18), 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    note = ws.max_row + 2
    ws.cell(row=note, column=1, value="Example rows - replace them with your project piping classes. Put each temperature "
                                      "and its rated pressure on its own line in the same cell (F / psig, or C / barg).").font = PLAIN
    return _save(wb)


TEMPLATES = {
    "risk-matrix.xlsx": ("HAZOP Studio - Risk Matrix template.xlsx", risk_matrix_template),
    "piping-spec.xlsx": ("HAZOP Studio - Piping Specification template.xlsx", piping_spec_template),
}


# ------------------------------------------------------------ optional docs for AI review

def _docx_text(data):
    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for row in t.rows:
            cells = list(dict.fromkeys(c.text.strip() for c in row.cells if c.text.strip()))
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _xlsx_text(data):
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append(f"[Sheet {ws.title}]")
        for row in ws.iter_rows(values_only=True):
            vals = [_text(v) for v in row if _text(v)]
            if vals:
                parts.append(" | ".join(vals))
    return "\n".join(parts)


def ai_attachment(name, data, render_page):
    """Turn an optional document into {'name', 'png'} or {'name', 'text'} for the AI review; None if unreadable."""
    ext = ext_of(name)
    try:
        if ext in (".docx",):
            return {"name": name, "text": _docx_text(data)[:MAX_TEXT_CHARS]}
        if ext in SPREADSHEET_EXT:
            return {"name": name, "text": _xlsx_text(data)[:MAX_TEXT_CHARS]}
        if ext == ".txt":
            return {"name": name, "text": data.decode("utf-8", errors="replace")[:MAX_TEXT_CHARS]}
        return {"name": name, "png": render_page(data, name)}
    except Exception:
        return None
