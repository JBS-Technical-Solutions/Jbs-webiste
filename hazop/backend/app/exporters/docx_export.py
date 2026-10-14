"""HAZOP worksheet export to Word, using the company template's worksheet table."""

import copy
import io

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Pt, RGBColor

from ..config import TEMPLATE_DOCX
from .. import risk
from ..risk import evaluate

HEADER_FILL = "C0E0FF"
FONT = "Arial"

COL_DEVIATION, COL_CAUSE, COL_CONSEQ, COL_CC, COL_SAFEGUARD, COL_TAG, COL_C, COL_L, COL_RISK, COL_REC = range(10)


def _fill(color):
    return color.lstrip("#").upper()


def set_shading(tc, fill):
    tc_pr = tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)


def set_vmerge(tc, restart):
    tc_pr = tc.get_or_add_tcPr()
    vmerge = OxmlElement("w:vMerge")
    if restart:
        vmerge.set(qn("w:val"), "restart")
    span = tc_pr.find(qn("w:gridSpan"))
    anchor = span if span is not None else tc_pr.find(qn("w:tcW"))
    if anchor is not None:
        anchor.addnext(vmerge)
    else:
        tc_pr.insert(0, vmerge)


def fill_cell(tc, lines, bold=False):
    """Replace the cell text with one paragraph per line, keeping template formatting."""
    paragraphs = tc.findall(qn("w:p"))
    proto = paragraphs[0]
    for p in paragraphs:
        tc.remove(p)
    p_rpr = proto.find(qn("w:pPr") + "/" + qn("w:rPr"))
    for line in lines or [""]:
        p = copy.deepcopy(proto)
        for r in p.findall(qn("w:r")):
            p.remove(r)
        if line:
            r = OxmlElement("w:r")
            if p_rpr is not None:
                rpr = copy.deepcopy(p_rpr)
                if bold:
                    rpr.append(OxmlElement("w:b"))
                r.append(rpr)
            t = OxmlElement("w:t")
            t.set(qn("xml:space"), "preserve")
            t.text = str(line)
            r.append(t)
            p.append(r)
        tc.append(p)


def scale_table_widths(tbl, factor):
    for col in tbl.iter(qn("w:gridCol")):
        col.set(qn("w:w"), str(round(int(col.get(qn("w:w"))) * factor)))
    for tag in ("w:tcW", "w:wAfter"):
        for el in tbl.iter(qn(tag)):
            if el.get(qn("w:type")) == "dxa":
                el.set(qn("w:w"), str(round(int(el.get(qn("w:w"))) * factor)))


def set_table_borders(table):
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    table._tbl.tblPr.append(borders)


def set_run_font(run, size, bold=False, color=None):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:cs"), FONT)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def write_cell(cell, text, size=8, bold=False, fill=None):
    cell.text = ""
    set_run_font(cell.paragraphs[0].add_run(str(text)), size, bold)
    if fill:
        set_shading(cell._tc, fill)


class WorksheetBuilder:
    def __init__(self):
        self.doc = Document(str(TEMPLATE_DOCX))
        body = self.doc.element.body
        tmpl_tbl = body.find(qn("w:tbl"))
        rows = tmpl_tbl.findall(qn("w:tr"))
        self.row_proto = copy.deepcopy(rows[2])
        self.table_proto = copy.deepcopy(tmpl_tbl)
        for tr in self.table_proto.findall(qn("w:tr"))[2:]:
            self.table_proto.remove(tr)

        section = self.doc.sections[0]
        usable = Emu(section.page_width - section.left_margin - section.right_margin)
        grid_total = sum(int(c.get(qn("w:w"))) for c in tmpl_tbl.iter(qn("w:gridCol")))
        scale_table_widths(self.table_proto, usable.twips / grid_total)
        scale_table_widths(self.row_proto, usable.twips / grid_total)
        self.usable_twips = usable.twips

        for el in list(body):
            if el.tag != qn("w:sectPr"):
                body.remove(el)
        self.sect_pr = body.find(qn("w:sectPr"))

    def heading(self, text, size=20):
        p = self.doc.add_paragraph()
        set_run_font(p.add_run(text), size, bold=True)

    def label_line(self, label, value, size=9):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        set_run_font(p.add_run(f"{label}: "), size, bold=True)
        set_run_font(p.add_run(value), size)

    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def info_table(self, rows, label_ratio=0.15):
        table = self.doc.add_table(rows=len(rows), cols=2)
        table.autofit = False
        set_table_borders(table)
        label_w = int(self.usable_twips * label_ratio)
        for i, (label, value) in enumerate(rows):
            a, b = table.rows[i].cells
            write_cell(a, label, bold=True, fill=HEADER_FILL)
            write_cell(b, value)
            a.width = Pt(label_w / 20)
            b.width = Pt((self.usable_twips - label_w) / 20)
        self.doc.add_paragraph()

    def grid_table(self, header, rows, widths, fills=None):
        table = self.doc.add_table(rows=len(rows) + 1, cols=len(header))
        table.autofit = False
        set_table_borders(table)
        total = sum(widths)
        for j, text in enumerate(header):
            cell = table.rows[0].cells[j]
            write_cell(cell, text, bold=True, fill=HEADER_FILL)
            cell.width = Pt(self.usable_twips * widths[j] / total / 20)
        for i, row in enumerate(rows, start=1):
            for j, text in enumerate(row):
                cell = table.rows[i].cells[j]
                write_cell(cell, text, fill=(fills or {}).get((i - 1, j)))
                cell.width = Pt(self.usable_twips * widths[j] / total / 20)
        self.doc.add_paragraph()

    def worksheet_table(self, node, rec_texts, seen_recs):
        tbl = copy.deepcopy(self.table_proto)
        self.sect_pr.addprevious(tbl)
        for dev in node["deviations"]:
            for i, row in enumerate(dev["rows"]):
                tr = copy.deepcopy(self.row_proto)
                tcs = tr.findall(qn("w:tc"))
                row_risk = evaluate(row["c"], row["l"])
                rec_lines = []
                for rid in row.get("recommendations", []):
                    if rid in seen_recs:
                        rec_lines.append(f"{rid}: Refer above.")
                    else:
                        rec_lines.append(f"{rid}: {rec_texts.get(rid, '')}")
                        seen_recs.add(rid)
                fill_cell(tcs[COL_DEVIATION], [dev["deviation"]] if i == 0 else [], bold=True)
                set_vmerge(tcs[COL_DEVIATION], restart=(i == 0))
                fill_cell(tcs[COL_CAUSE], [f"{i + 1}. {row['cause']}"])
                fill_cell(tcs[COL_CONSEQ], [f"{n}. {c}" for n, c in enumerate(row.get("consequences", []), 1)])
                fill_cell(tcs[COL_CC], [row.get("cc", "")])
                fill_cell(tcs[COL_SAFEGUARD], [f"{n}. {s}" for n, s in enumerate(row.get("safeguards", []), 1)])
                fill_cell(tcs[COL_TAG], row.get("tags", []))
                fill_cell(tcs[COL_C], [str(row["c"])])
                fill_cell(tcs[COL_L], [str(row["l"])])
                fill_cell(tcs[COL_RISK], [str(row_risk["score"]), row_risk["code"]], bold=True)
                set_shading(tcs[COL_RISK], _fill(row_risk["color"]))
                fill_cell(tcs[COL_REC], rec_lines)
                tbl.append(tr)


def build_docx(result):
    study = result.get("study", {})
    recs = result.get("recommendations", {})
    b = WorksheetBuilder()
    b.doc.core_properties.title = f"{study.get('title', 'HAZOP Worksheets')} - {study.get('drawing', '')}".strip(" -")

    b.heading(study.get("title") or "HAZOP Worksheets")
    for label, key in (("Project", "project"), ("Facility", "facility"), ("Company", "client"),
                       ("Drawing", "drawing"), ("Revision", "revision"), ("Study date", "date"),
                       ("HAZOP team", "team"), ("Reference documents", "reference_documents")):
        if study.get(key):
            b.label_line(label, str(study[key]))
    b.doc.add_paragraph()

    seen, rec_refs = set(), {}
    for idx, node in enumerate(result.get("nodes", [])):
        if idx:
            b.page_break()
        b.heading(f"Node {node['id']}: {node['title']}", size=12)
        b.info_table([
            ("Node boundary", node.get("boundary", "")),
            ("Design intent", node.get("design_intent", "")),
            ("Design conditions", node.get("design_conditions", "")),
        ])
        b.worksheet_table(node, recs, seen)
        for dev in node["deviations"]:
            for row in dev["rows"]:
                for rid in row.get("recommendations", []):
                    info = rec_refs.setdefault(rid, {"refs": [], "worst": None})
                    info["refs"].append(dev["deviation"].split(" ")[0])
                    info["worst"] = risk.worse(info["worst"], evaluate(row["c"], row["l"]))

    b.page_break()
    b.heading("Recommendations Summary", size=12)
    rows, fills = [], {}
    for i, (rid, text) in enumerate(recs.items()):
        info = rec_refs.get(rid, {"refs": [], "worst": None})
        refs = ", ".join(dict.fromkeys(info["refs"]))
        worst = info["worst"]
        if worst:
            rows.append([rid, text, refs, f"{worst['score']} ({worst['name']})", "", "Open"])
            fills[(i, 3)] = _fill(worst["color"])
        else:
            rows.append([rid, text, refs, "-", "", "Open"])
    b.grid_table(["No.", "Recommendation", "Deviation Ref.", "Highest Risk", "Action By", "Status"],
                 rows, [5, 60, 10, 9, 9, 7], fills)

    m = risk.current()
    severity, likelihood = m["severity"], m["likelihood"]
    heading = "Risk matrix" if m.get("scores") else "Risk Matrix (Risk = C x L)"
    b.heading(f"{heading} - {m['name']}" + (f", {m['source']}" if m.get("source") else ""), size=12)
    matrix_rows, matrix_fills = [], {}
    for r, sev in enumerate(reversed(severity)):
        row = [f"C{sev['level']} - {sev['name']}"]
        for j, lk in enumerate(likelihood, start=1):
            cell_risk = evaluate(sev["level"], lk["level"])
            row.append(f"{cell_risk['score']} ({cell_risk['code']})")
            matrix_fills[(r, j)] = _fill(cell_risk["color"])
        matrix_rows.append(row)
    b.grid_table(["Consequence \\ Likelihood"] + [f"L{lk['level']} - {lk['name']}" for lk in likelihood],
                 matrix_rows, [16] + [12] * len(likelihood), matrix_fills)
    b.grid_table(["C", "Consequence", "Description"],
                 [[s["level"], s["name"], s["desc"]] for s in severity], [4, 12, 70])
    b.grid_table(["L", "Likelihood", "Description"],
                 [[lk["level"], lk["name"], lk["desc"]] for lk in likelihood], [4, 12, 70])
    b.grid_table(["Risk", "Ranking"], [[bd["code"], bd["name"]] for bd in m["bands"]],
                 [6, 80], {(i, 0): _fill(bd["color"]) for i, bd in enumerate(m["bands"])})

    if result.get("assumptions"):
        b.heading("Assumptions and Basis", size=12)
        for a in result["assumptions"]:
            set_run_font(b.doc.add_paragraph().add_run(f"\u2022 {a}"), 9)

    buf = io.BytesIO()
    b.doc.save(buf)
    return buf.getvalue()
