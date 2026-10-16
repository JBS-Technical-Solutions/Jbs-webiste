"""HAZOP worksheet export to Excel."""

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .. import risk
from ..risk import evaluate

HEADER_FILL = PatternFill("solid", fgColor="C0E0FF")
NODE_FILL = PatternFill("solid", fgColor="1E3A5F")
INFO_FILL = PatternFill("solid", fgColor="EEF4FB")
THIN = Side(style="thin", color="7F8C9D")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(wrap_text=True, vertical="center", horizontal="center")

WORKSHEET_COLS = [("Deviation", 22), ("Causes", 34), ("Consequences", 48), ("CC", 6), ("Safeguards", 44),
                  ("Tag #", 18), ("C", 5), ("L", 5), ("Risk", 9), ("Recommendations", 60)]


def _fill(color):
    return PatternFill("solid", fgColor=color.lstrip("#").upper())


def _numbered(items):
    return "\n".join(f"{i}. {x}" for i, x in enumerate(items, 1))


def _header(ws, row, titles):
    for j, title in enumerate(titles, 1):
        c = ws.cell(row=row, column=j, value=title)
        c.font = Font(name="Arial", bold=True, size=10)
        c.fill = HEADER_FILL
        c.alignment = CENTER
        c.border = BORDER


def _worksheet(wb, result):
    ws = wb.active
    ws.title = "HAZOP Worksheet"
    study = result.get("study", {})
    ncol = len(WORKSHEET_COLS)
    for j, (_, width) in enumerate(WORKSHEET_COLS, 1):
        ws.column_dimensions[get_column_letter(j)].width = width

    ws.cell(row=1, column=1, value=study.get("title") or "HAZOP Worksheets").font = Font(name="Arial", bold=True, size=16)
    r = 2
    for label, key in (("Project", "project"), ("Facility", "facility"), ("Company", "client"), ("Drawing", "drawing"),
                       ("Revision", "revision"), ("Study date", "date"), ("HAZOP team", "team")):
        if study.get(key):
            ws.cell(row=r, column=1, value=label).font = Font(name="Arial", bold=True, size=10)
            ws.cell(row=r, column=2, value=str(study[key])).font = Font(name="Arial", size=10)
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=ncol)
            r += 1
    r += 1

    recs = result.get("recommendations", {})
    seen = set()
    for node in result.get("nodes", []):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncol)
        c = ws.cell(row=r, column=1, value=f"Node {node['id']}: {node['title']}")
        c.font = Font(name="Arial", bold=True, size=12, color="FFFFFF")
        c.fill = NODE_FILL
        c.alignment = Alignment(vertical="center")
        ws.row_dimensions[r].height = 22
        r += 1
        for label, key in (("Node description", "boundary"), ("Design intent", "design_intent"),
                           ("Design conditions", "design_conditions")):
            a = ws.cell(row=r, column=1, value=label)
            a.font = Font(name="Arial", bold=True, size=9)
            a.fill = INFO_FILL
            a.border = BORDER
            a.alignment = WRAP
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=ncol)
            b = ws.cell(row=r, column=2, value=node.get(key, ""))
            b.font = Font(name="Arial", size=9)
            b.alignment = WRAP
            for j in range(2, ncol + 1):
                ws.cell(row=r, column=j).border = BORDER
            ws.row_dimensions[r].height = max(15, 15 * (len(node.get(key, "")) // 220 + 1))
            r += 1
        _header(ws, r, [t for t, _ in WORKSHEET_COLS])
        r += 1
        for dev in node["deviations"]:
            start = r
            for i, row in enumerate(dev["rows"]):
                row_risk = evaluate(row["c"], row["l"])
                rec_lines = []
                for rid in row.get("recommendations", []):
                    rec_lines.append(f"{rid}: Refer above." if rid in seen else f"{rid}: {recs.get(rid, '')}")
                    seen.add(rid)
                values = [dev["deviation"] if i == 0 else None, f"{i + 1}. {row['cause']}",
                          _numbered(row.get("consequences", [])), row.get("cc", ""),
                          _numbered(row.get("safeguards", [])), "\n".join(row.get("tags", [])),
                          int(row["c"]), int(row["l"]), f"{row_risk['score']} {row_risk['code']}", "\n".join(rec_lines)]
                for j, v in enumerate(values, 1):
                    c = ws.cell(row=r, column=j, value=v)
                    c.font = Font(name="Arial", size=9, bold=(j in (1, 9)))
                    c.alignment = CENTER if j in (4, 7, 8, 9) else WRAP
                    c.border = BORDER
                ws.cell(row=r, column=9).fill = _fill(row_risk["color"])
                longest = max(len(str(v or "")) / max(w * 1.1, 1) for v, (_, w) in zip(values, WORKSHEET_COLS))
                lines = max(len(str(values[j] or "").split("\n")) for j in range(ncol))
                ws.row_dimensions[r].height = min(400, 13 * max(lines, int(longest) + 1, 2))
                r += 1
            if r - start > 1:
                ws.merge_cells(start_row=start, start_column=1, end_row=r - 1, end_column=1)
        r += 1
    ws.freeze_panes = "A2"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def _recommendations(wb, result):
    ws = wb.create_sheet("Recommendations")
    widths = [8, 90, 18, 16, 16, 10]
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    _header(ws, 1, ["No.", "Recommendation", "Deviation Ref.", "Highest Risk", "Action By", "Status"])
    refs = {}
    for node in result.get("nodes", []):
        for dev in node["deviations"]:
            for row in dev["rows"]:
                for rid in row.get("recommendations", []):
                    info = refs.setdefault(rid, {"refs": [], "worst": None})
                    info["refs"].append(dev["deviation"].split(" ")[0])
                    info["worst"] = risk.worse(info["worst"], evaluate(row["c"], row["l"]))
    for i, (rid, text) in enumerate(result.get("recommendations", {}).items(), 2):
        info = refs.get(rid, {"refs": [], "worst": None})
        worst = info["worst"]
        vals = [rid, text, ", ".join(dict.fromkeys(info["refs"])),
                f"{worst['score']} ({worst['name']})" if worst else "-", "", "Open"]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=i, column=j, value=v)
            c.font = Font(name="Arial", size=9, bold=(j == 1))
            c.alignment = WRAP
            c.border = BORDER
        if worst:
            ws.cell(row=i, column=4).fill = _fill(worst["color"])
        ws.row_dimensions[i].height = 13 * (len(text) // 95 + 1)
    ws.freeze_panes = "A2"


def _matrix(wb):
    m = risk.current()
    severity, likelihood = m["severity"], m["likelihood"]
    last = max(len(likelihood) + 1, 6)
    ws = wb.create_sheet("Risk Matrix")
    ws.column_dimensions["A"].width = 26
    for j in range(2, len(likelihood) + 2):
        ws.column_dimensions[get_column_letter(j)].width = 16
    title = ws.cell(row=1, column=1, value=f"{m['name']}" + (f" ({m['source']})" if m.get("source") else ""))
    title.font = Font(name="Arial", bold=True, size=11)
    _header(ws, 2, ["Consequence \\ Likelihood"] + [f"L{l['level']} - {l['name']}" for l in likelihood])
    for i, sev in enumerate(reversed(severity), 3):
        a = ws.cell(row=i, column=1, value=f"C{sev['level']} - {sev['name']}")
        a.font = Font(name="Arial", bold=True, size=10)
        a.border = BORDER
        for j, lk in enumerate(likelihood, 2):
            cell_risk = evaluate(sev["level"], lk["level"])
            c = ws.cell(row=i, column=j, value=f"{cell_risk['score']} ({cell_risk['code']})")
            c.fill = _fill(cell_risk["color"])
            c.alignment = CENTER
            c.border = BORDER
            c.font = Font(name="Arial", bold=True, size=10)
        ws.row_dimensions[i].height = 26
    r = len(severity) + 4
    _header(ws, r, ["Risk", "Ranking"])
    for k, band in enumerate(m["bands"], 1):
        a = ws.cell(row=r + k, column=1, value=band["code"])
        a.fill = _fill(band["color"])
        a.font = Font(name="Arial", bold=True, size=9)
        b = ws.cell(row=r + k, column=2, value=band["name"])
        b.font = Font(name="Arial", size=9)
        a.border = b.border = BORDER
    r += len(m["bands"]) + 2
    for heading, items in (("Consequence", severity), ("Likelihood", likelihood)):
        _header(ws, r, ["Level", heading, "Description"])
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=last)
        for k, it in enumerate(items, 1):
            for j, v in enumerate([it["level"], it["name"], it["desc"]], 1):
                c = ws.cell(row=r + k, column=j, value=v)
                c.font = Font(name="Arial", size=9)
                c.border = BORDER
            ws.merge_cells(start_row=r + k, start_column=3, end_row=r + k, end_column=last)
        r += len(items) + 2


def _extracted(wb, ext):
    if not ext:
        return
    sheets = [
        ("Equipment", ["Tag", "Description", "Type", "API 14C", "Design", "Off-page"],
         lambda e: [e["tag"], e.get("description", ""), e.get("kind", ""), e.get("api14c") or "",
                    ", ".join(f"{k}={v}" for k, v in (e.get("design") or {}).items()), "Yes" if e.get("offpage") else ""],
         ext.get("equipment", [])),
        ("Instruments", ["Tag", "Description", "Category", "Details"],
         lambda i: [i["tag"], i.get("description", ""), i.get("category", ""),
                    ", ".join(f"{k}={i[k]}" for k in ("set_point", "set_psig", "fail_action", "voting", "remark", "inferred")
                              if i.get(k) not in (None, ""))],
         ext.get("instruments", [])),
        ("Lines", ["Line", "Service", "Spec", "Rating", "Material", "Temp range (F)"],
         lambda l: [l["tag"], l.get("service_name") or l.get("service"), l.get("spec"), l.get("rating") or "",
                    l.get("material") or "", f"{l.get('min_temp_f')} to {l.get('max_temp_f')}" if l.get("spec_known") else ""],
         ext.get("lines", [])),
        ("Valves", ["Code", "Type", "Size", "Attributes"],
         lambda v: [v["code"], v.get("type", ""), v.get("size") or "", ", ".join(v.get("attributes", []))],
         ext.get("valves", [])),
    ]
    for title, header, fn, items in sheets:
        if any(isinstance(it, dict) and "page" in it for it in items):
            header = ["Sheet", *header]
            row_fn = fn
            fn = lambda it, row_fn=row_fn: [it.get("page", 0) + 1, *row_fn(it)]
        ws = wb.create_sheet(title)
        _header(ws, 1, header)
        for i, item in enumerate(items, 2):
            for j, v in enumerate(fn(item), 1):
                c = ws.cell(row=i, column=j, value=v)
                c.font = Font(name="Arial", size=9)
                c.border = BORDER
                c.alignment = WRAP
        for j in range(1, len(header) + 1):
            ws.column_dimensions[get_column_letter(j)].width = 18 if j == 1 else 34
        ws.freeze_panes = "A2"


def build_xlsx(result, extraction=None):
    wb = Workbook()
    _worksheet(wb, result)
    _recommendations(wb, result)
    _matrix(wb)
    _extracted(wb, extraction)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
