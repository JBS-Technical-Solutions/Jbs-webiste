"""Free sign-up: a person confirms their email with a 6-digit code and gets a token that unlocks the analysis.

Sign-ups live in SQLite (data/signups.db); tokens and codes are stored only as hashes.
"""

import hashlib
import hmac
import io
import json
import logging
import re
import secrets
import smtplib
import sqlite3
import ssl
import threading
import time
from contextlib import contextmanager
from datetime import datetime, timezone
from email.message import EmailMessage
from email.utils import formataddr

import dns.exception
import dns.resolver
from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from .config import DATA_DIR, SMTP_FROM, SMTP_HOST, SMTP_PASSWORD, SMTP_PORT, SMTP_USER

log = logging.getLogger("hazop.signup")

DB_PATH = DATA_DIR / "signups.db"
CODE_TTL_S = 10 * 60
RESEND_AFTER_S = 60
MAX_CODES_PER_HOUR = 5
MAX_ATTEMPTS = 5

EMAIL_RE = re.compile(
    r"^[A-Za-z0-9!#$%&'*+/=?^_`{|}~-]+(?:\.[A-Za-z0-9!#$%&'*+/=?^_`{|}~-]+)*"
    r"@(?:[A-Za-z0-9](?:[A-Za-z0-9-]{0,61}[A-Za-z0-9])?\.)+[A-Za-z]{2,63}$"
)
NAME_RE = re.compile(r"^[^\W\d_](?:[^\W\d_]|[ .'\-])*$")
DISPOSABLE_DOMAINS = {
    "10minutemail.com", "burnermail.io", "dispostable.com", "emailondeck.com", "fakeinbox.com", "getnada.com",
    "guerrillamail.com", "guerrillamail.net", "maildrop.cc", "mailinator.com", "mailnesia.com", "mintemail.com",
    "mohmal.com", "sharklasers.com", "temp-mail.org", "tempail.com", "tempmail.com", "tempmail.dev", "tempr.email",
    "throwawaymail.com", "trashmail.com", "yopmail.com",
}
FIELDS = {"first_name": ("First name", 60), "last_name": ("Last name", 60), "company": ("Company name", 120),
          "role": ("Role", 80)}

_lock = threading.Lock()
_mx_cache: dict[str, bool] = {}


@contextmanager
def _db():
    conn = sqlite3.connect(DB_PATH, timeout=10)
    conn.row_factory = sqlite3.Row
    try:
        with conn:
            yield conn
    finally:
        conn.close()


def init():
    with _db() as c:
        c.executescript("""
            CREATE TABLE IF NOT EXISTS users (
                email TEXT PRIMARY KEY, first_name TEXT NOT NULL, last_name TEXT NOT NULL, company TEXT NOT NULL,
                role TEXT NOT NULL, created_at TEXT NOT NULL, last_login_at TEXT NOT NULL);
            CREATE TABLE IF NOT EXISTS codes (
                email TEXT PRIMARY KEY, code_hash TEXT NOT NULL, details TEXT, expires REAL NOT NULL,
                attempts INTEGER NOT NULL DEFAULT 0, sent TEXT NOT NULL DEFAULT '[]');
            CREATE TABLE IF NOT EXISTS tokens (
                token_hash TEXT PRIMARY KEY, email TEXT NOT NULL, created_at TEXT NOT NULL);
        """)


def _now_iso():
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def _sha(value: str) -> str:
    return hashlib.sha256(value.encode()).hexdigest()


def _user(row) -> dict:
    return {k: row[k] for k in ("first_name", "last_name", "email", "company", "role")}


def _email(value) -> str:
    email = str(value or "").strip().lower()
    if not email:
        raise HTTPException(400, "Enter your email address.")
    local, _, domain = email.rpartition("@")
    if len(email) > 254 or len(local) > 64 or not EMAIL_RE.match(email):
        raise HTTPException(400, "This email address is not valid. Check it for typos.")
    if domain in DISPOSABLE_DOMAINS:
        raise HTTPException(400, "Temporary email addresses can't be used. Enter your work email.")
    return email


def _receives_mail(domain: str) -> bool:
    """True unless DNS says the domain does not exist or accepts no mail; DNS trouble is left to the code email."""
    if domain in _mx_cache:
        return _mx_cache[domain]
    resolver = dns.resolver.Resolver()
    resolver.lifetime = 5
    ok = False
    for rtype in ("MX", "A", "AAAA"):
        try:
            answers = resolver.resolve(domain, rtype)
        except dns.resolver.NXDOMAIN:
            break
        except dns.resolver.NoAnswer:
            continue
        except dns.exception.DNSException:
            return True
        ok = not (rtype == "MX" and all(str(r.exchange) == "." for r in answers))
        break
    _mx_cache[domain] = ok
    return ok


def _details(payload: dict) -> dict:
    out = {}
    for key, (label, max_len) in FIELDS.items():
        value = re.sub(r"\s+", " ", str(payload.get(key) or "")).strip()
        if not value:
            raise HTTPException(400, f"Enter your {label.lower()}.")
        if len(value) > max_len:
            raise HTTPException(400, f"{label} is too long (at most {max_len} characters).")
        if key in ("first_name", "last_name") and not NAME_RE.match(value):
            raise HTTPException(400, f"{label} can only contain letters, spaces, dots, hyphens and apostrophes.")
        out[key] = value
    return out


def request_code(payload: dict) -> dict:
    """Send a sign-in code. payload: {email, mode: "signup"|"login", first_name, last_name, company, role}."""
    login = payload.get("mode") == "login"
    email = _email(payload.get("email"))
    details = None if login else _details(payload)
    with _db() as c:
        user = c.execute("SELECT first_name FROM users WHERE email = ?", (email,)).fetchone()
    if login and not user:
        raise HTTPException(404, "There is no account for this email yet. Sign up instead, it's free.")
    if not user and not _receives_mail(email.rpartition("@")[2]):
        raise HTTPException(400, f"The domain '{email.rpartition('@')[2]}' doesn't receive email. Check the address for typos.")

    now = time.time()
    code = f"{secrets.randbelow(1_000_000):06d}"
    with _lock, _db() as c:
        row = c.execute("SELECT sent FROM codes WHERE email = ?", (email,)).fetchone()
        sent = [t for t in json.loads(row["sent"]) if now - t < 3600] if row else []
        if sent and now - sent[-1] < RESEND_AFTER_S:
            raise HTTPException(429, f"Please wait {int(RESEND_AFTER_S - (now - sent[-1])) + 1} seconds before asking for another code.")
        if len(sent) >= MAX_CODES_PER_HOUR:
            raise HTTPException(429, "Too many codes were sent to this email. Please try again in an hour.")
        c.execute("INSERT OR REPLACE INTO codes (email, code_hash, details, expires, attempts, sent) VALUES (?, ?, ?, ?, 0, ?)",
                  (email, _sha(f"{email}:{code}"), json.dumps(details) if details else None, now + CODE_TTL_S,
                   json.dumps(sent + [now])))
    _send_code(email, code, (details or {}).get("first_name") or (user["first_name"] if user else ""))
    return {"email": email, "expires_in": CODE_TTL_S, "resend_after": RESEND_AFTER_S}


def verify_code(payload: dict) -> dict:
    """Check a code; returns {token, user, new}."""
    email = _email(payload.get("email"))
    code = re.sub(r"\D", "", str(payload.get("code") or ""))
    now = time.time()
    problem = None
    with _lock, _db() as c:
        row = c.execute("SELECT * FROM codes WHERE email = ?", (email,)).fetchone()
        if not row or row["expires"] < now or row["code_hash"] == "":
            problem = (410, "This code has expired. Ask for a new one.")
        elif row["attempts"] >= MAX_ATTEMPTS:
            problem = (429, "Too many wrong codes. Ask for a new one.")
        elif not hmac.compare_digest(row["code_hash"], _sha(f"{email}:{code}")):
            c.execute("UPDATE codes SET attempts = attempts + 1 WHERE email = ?", (email,))
            left = MAX_ATTEMPTS - row["attempts"] - 1
            problem = (400, f"That code is not correct. {left} attempt{'s' if left != 1 else ''} left." if left
                       else "That code is not correct. Ask for a new one.")
        else:
            c.execute("UPDATE codes SET code_hash = '', details = NULL WHERE email = ?", (email,))
            stamp = _now_iso()
            details = json.loads(row["details"]) if row["details"] else None
            existing = c.execute("SELECT 1 FROM users WHERE email = ?", (email,)).fetchone()
            if details:
                c.execute(
                    "INSERT INTO users (email, first_name, last_name, company, role, created_at, last_login_at) "
                    "VALUES (?, ?, ?, ?, ?, ?, ?) ON CONFLICT(email) DO UPDATE SET first_name = excluded.first_name, "
                    "last_name = excluded.last_name, company = excluded.company, role = excluded.role, "
                    "last_login_at = excluded.last_login_at",
                    (email, details["first_name"], details["last_name"], details["company"], details["role"], stamp, stamp))
            elif existing:
                c.execute("UPDATE users SET last_login_at = ? WHERE email = ?", (stamp, email))
            user = c.execute("SELECT * FROM users WHERE email = ?", (email,)).fetchone()
            if not user:
                problem = (404, "There is no account for this email yet. Sign up instead, it's free.")
            else:
                token = secrets.token_urlsafe(32)
                c.execute("INSERT INTO tokens (token_hash, email, created_at) VALUES (?, ?, ?)", (_sha(token), email, stamp))
    if problem:
        raise HTTPException(*problem)
    return {"token": token, "user": _user(user), "new": not existing}


def user_for_token(token: str | None) -> dict | None:
    if not token:
        return None
    with _db() as c:
        row = c.execute("SELECT u.* FROM tokens t JOIN users u ON u.email = t.email WHERE t.token_hash = ?",
                        (_sha(token),)).fetchone()
    return _user(row) if row else None


def logout(token: str | None):
    if token:
        with _db() as c:
            c.execute("DELETE FROM tokens WHERE token_hash = ?", (_sha(token),))


def export_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sign-ups"
    headers = ["First name", "Last name", "Email", "Company", "Role", "Signed up (UTC)", "Last sign-in (UTC)"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E79")
    with _db() as c:
        for r in c.execute("SELECT * FROM users ORDER BY created_at DESC"):
            ws.append([r["first_name"], r["last_name"], r["email"], r["company"], r["role"],
                       r["created_at"].replace("T", " ")[:19], r["last_login_at"].replace("T", " ")[:19]])
    for col, width in zip("ABCDEFG", (16, 18, 32, 28, 26, 20, 20)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _send_code(email: str, code: str, name: str):
    if not SMTP_HOST:
        log.warning("SMTP_HOST is not set, so no email was sent. Sign-in code for %s: %s", email, code)
        return
    greeting = f"Hi {name}," if name else "Hi,"
    msg = EmailMessage()
    msg["Subject"] = f"{code} is your HAZOP Studio code"
    msg["From"] = formataddr(("HAZOP Studio", SMTP_FROM))
    msg["To"] = email
    msg.set_content(f"{greeting}\n\nYour HAZOP Studio code is {code}\n\nIt expires in {CODE_TTL_S // 60} minutes. "
                    "If you didn't ask for it, you can ignore this email.\n\nJBS Technical Solutions")
    msg.add_alternative(f"""\
<div style="font-family:Segoe UI,Arial,sans-serif;max-width:480px;margin:auto;color:#0f172a">
  <p>{greeting}</p>
  <p>Your HAZOP Studio code is</p>
  <p style="font-size:32px;font-weight:700;letter-spacing:8px;margin:16px 0">{code}</p>
  <p style="color:#475569;font-size:14px">It expires in {CODE_TTL_S // 60} minutes.
  If you didn't ask for it, you can ignore this email.</p>
  <p style="color:#475569;font-size:14px">JBS Technical Solutions</p>
</div>""", subtype="html")
    try:
        if SMTP_PORT == 465:
            server = smtplib.SMTP_SSL(SMTP_HOST, SMTP_PORT, timeout=20, context=ssl.create_default_context())
        else:
            server = smtplib.SMTP(SMTP_HOST, SMTP_PORT, timeout=20)
            server.starttls(context=ssl.create_default_context())
        with server:
            if SMTP_USER:
                server.login(SMTP_USER, SMTP_PASSWORD)
            server.send_message(msg)
    except (smtplib.SMTPException, OSError):
        log.exception("Could not send the sign-in code to %s", email)
        raise HTTPException(502, "We couldn't send the code email right now. Please try again in a few minutes.")
